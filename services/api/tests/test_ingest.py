from __future__ import annotations

import io
import shutil
import socket
import zipfile

import pytest

from kila.settings import REPO_ROOT

SEED = REPO_ROOT / "seed"


def _ingest(c, name: str, data: bytes, bucket: str = "uploads", **params) -> dict:
    from kila.ingest import service

    oid = c.post(f"/storage/{bucket}/upload", files={"file": (name, data)}).json()["object_id"]
    assert c.post(f"/ingest/{oid}", params=params).status_code == 200
    service.wait_idle(300)
    r = c.get(f"/ingest/{oid}").json()
    assert r["status"] == "done", r
    return r


def test_text_pdf_uses_text_layer_and_page_images(client_for):
    eng = client_for("engineer")
    a = _ingest(eng, "SOP-MNT-010.pdf", (SEED / "kb/SOP-MNT-010.pdf").read_bytes())["attachment"]
    assert a["kind"] == "pdf_text"
    p = a["pages"][0]
    assert p["method"] == "text_layer" and "set pressure" in p["text"] and "\r" not in p["text"]
    assert p["image_object_id"] and eng.get(f"/storage/object/{p['image_object_id']}").status_code == 200


def test_scanned_pdf_is_deskewed_and_ocrd(client_for):
    eng = client_for("engineer")
    r = _ingest(eng, "IR-2026-0147_scan.pdf", (SEED / "reports/IR-2026-0147_scan.pdf").read_bytes())
    a = r["attachment"]
    p = a["pages"][0]
    assert a["kind"] == "pdf_scanned" and p["method"] == "ocr"
    assert p["ocr_conf"] > 0.9 and abs(p["skew_deg"]) > 0.5  # the synthetic scan is rotated ~1.6 degrees
    assert "IR-2026-0147" in p["text"]
    # table rows are rebuilt from separate cell detections
    assert any(line.startswith("C3") and "0.425" in line for line in p["text"].splitlines())
    assert p["lines"] and len(p["lines"][0]["box"]) == 4
    events = [e["event_type"] for e in eng.get("/ledger/events?limit=50").json()]
    assert "ingest.queued" in events and "ingest.done" in events


def test_regression_orientation_classifier_does_not_garble_lines(client_for):
    """With PP-OCR's 180-degree line classifier on, this exact line came back as 'S(et d - ( d - ( g)' at
    63% confidence. Documents are upright after deskew, so the classifier is off (config/ingest.yaml)."""
    a = _ingest(client_for("engineer"), "IR-2026-0163_scan.pdf",
                (SEED / "reports/IR-2026-0163_scan.pdf").read_bytes())["attachment"]
    p = a["pages"][0]
    assert "as-received pop 3.71 bar(g)" in p["text"]
    assert min(ln["conf"] for ln in p["lines"]) > 0.9


def test_spreadsheet_summary_and_rows(client_for):
    eng = client_for("engineer")
    a = _ingest(eng, "psv_register.xlsx", (SEED / "sheets/psv_register.xlsx").read_bytes())["attachment"]
    t = a["tables"][0]
    assert a["kind"] == "sheet" and t["n_rows"] == 5 and t["columns"][0] == "Tag"
    assert t["rows_json"][0]["Tag"] == "PSV-1204" and "float" in t["dtypes"]["Set pressure (barg)"]
    assert "PSV-1204" in a["pages"][0]["text"]  # searchable text rendering


def test_csv(client_for):
    a = _ingest(client_for("engineer"), "t.csv", b"tag,value\nPV-101,12.5\nFT-2031,340\n")["attachment"]
    assert a["kind"] == "sheet" and a["tables"][0]["n_rows"] == 2


def test_docx(client_for):
    import docx

    d = docx.Document()
    d.add_paragraph("Vendor letter about PSV-1204 spares.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Item", "Qty"
    buf = io.BytesIO()
    d.save(buf)
    a = _ingest(client_for("engineer"), "letter.docx", buf.getvalue())["attachment"]
    assert a["kind"] == "docx" and "PSV-1204" in a["pages"][0]["text"] and "Item | Qty" in a["pages"][0]["text"]


def test_code_zip_lists_files_and_rejects_traversal(client_for):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("pkg/rates.py", "def f():\n    return 1\n")
        z.writestr("../evil.py", "import os\n")
        z.writestr("img.bin", b"\x00\x01")
    a = _ingest(client_for("engineer"), "code.zip", buf.getvalue())["attachment"]
    paths = [f["path"] for f in a["code"]]
    assert a["kind"] == "code" and "pkg/rates.py" in paths and "../evil.py" not in paths
    assert any("unsafe path" in w for w in a["warnings"])
    assert next(f for f in a["code"] if f["path"] == "img.bin")["text"] is None


def test_identical_content_is_not_processed_twice(client_for):
    eng = client_for("engineer")
    data = (SEED / "kb/SOP-OPS-001.pdf").read_bytes()
    first = _ingest(eng, "a.pdf", data)
    second = _ingest(eng, "b.pdf", data)
    assert second["duration_ms"] == 0.0 and second["attachment"]["name"] == "b.pdf"
    assert second["attachment"]["pages"][0]["text"] == first["attachment"]["pages"][0]["text"]
    assert eng.get("/ledger/events?event_type=ingest.cache_hit").json()


def test_ocr_never_touches_the_network(client_for, monkeypatch):
    """RapidOCR has download code; KILA must never reach it. Block all sockets and OCR a scan."""
    from kila.ingest import ocr

    ocr._engines.clear()  # force engine construction under the block

    real_connect, real_create = socket.socket.connect, socket.create_connection
    loopback = ("127.0.0.1", "::1", "localhost")  # asyncio's self-pipe on Windows uses loopback

    def guarded_connect(sock, address, *a, **kw):
        if isinstance(address, tuple) and address[0] not in loopback:
            raise AssertionError(f"network access attempted during OCR: {address}")
        return real_connect(sock, address, *a, **kw)

    def guarded_create(address, *a, **kw):
        if address[0] not in loopback:
            raise AssertionError(f"network access attempted during OCR: {address}")
        return real_create(address, *a, **kw)

    eng = client_for("engineer")
    monkeypatch.setattr(socket.socket, "connect", guarded_connect)
    monkeypatch.setattr(socket, "create_connection", guarded_create)
    a = _ingest(eng, "scan.pdf", (SEED / "reports/IR-2026-0163_scan.pdf").read_bytes())["attachment"]
    assert a["pages"][0]["method"] == "ocr" and "PSV-1204" in a["pages"][0]["text"]


def test_vision_second_opinion_on_low_confidence(tmp_path, monkeypatch, client_for):
    """Force the low-confidence path: the vision model (fake server) is consulted and disagreement flagged."""
    from kila.settings import get_settings, load_app_config

    cfg_dir = tmp_path / "cfg"
    shutil.copytree(REPO_ROOT / "config", cfg_dir)
    text = (cfg_dir / "ingest.yaml").read_text(encoding="utf-8").replace("low_confidence: 0.80", "low_confidence: 1.01")
    (cfg_dir / "ingest.yaml").write_text(text, encoding="utf-8")
    monkeypatch.setenv("KILA_CONFIG_DIR", str(cfg_dir))
    get_settings.cache_clear()
    load_app_config.cache_clear()
    a = _ingest(client_for("engineer"), "scan.pdf", (SEED / "reports/IR-2026-0155_scan.pdf").read_bytes())["attachment"]
    p = a["pages"][0]
    assert p["method"] == "ocr+vision" and p["vision"]["model"] == "qwen3.5:4b"
    assert p["vision"]["disagrees"] is True and any("disagree" in w for w in a["warnings"])
    assert "P-1201B" in p["text"]  # OCR text is kept; the vision reading never silently replaces it


def test_interrupted_jobs_are_requeued(app_env, client_for):
    from sqlmodel import Session

    from kila.db.models import Ingestion
    from kila.db.session import get_engine
    from kila.ingest import service

    eng = client_for("engineer")
    oid = eng.post("/storage/uploads/upload", files={"file": ("n.txt", b"notes about PV-101")}).json()["object_id"]
    with Session(get_engine()) as s:  # a job that was 'running' when the API died
        s.add(Ingestion(object_id=oid, sha256="x", pipeline=service.PIPELINE, status="running"))
        s.commit()
    assert service.requeue_interrupted() == 1
    service.wait_idle(60)
    assert eng.get(f"/ingest/{oid}").json()["status"] == "done"


def test_reviewer_can_read_but_not_touch_models_bucket(client_for):
    admin, rev = client_for("admin"), client_for("reviewer")
    oid = admin.post("/storage/models/upload", files={"file": ("m.txt", b"x")}).json()["object_id"]
    assert rev.post(f"/ingest/{oid}").status_code == 403


@pytest.mark.parametrize("text,lang", [("PSV test due", "en"), ("रिसाव की जाँच करें", "hi"), ("ಒತ್ತಡ ಪರೀಕ್ಷೆ", "kn")])
def test_language_detection(text, lang):
    from kila.ingest.envelope import detect_language

    assert detect_language(text) == lang
